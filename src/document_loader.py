"""Загрузка документов из файлов и директорий: PDF, DOCX, TXT/MD."""

from dataclasses import dataclass
from pathlib import Path
from typing import List, Union

from docx import Document as DocxDocument
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
TEXT_ENCODINGS = ("utf-8-sig", "cp1251")  # utf-8-sig снимает Byte Order Mark


@dataclass
class LoadedDocument:
    """Документ с привязкой к источнику - нужна, чтобы ассистент мог сослаться на файл."""
    source: str  # Имя файла, как ссылка на источник
    text: str


def _read_pdf(path: Path) -> str:
    """Собирает текст со всех страниц PDF, пропуская страницы без текстового слоя."""
    reader = PdfReader(path)
    pages = [page.extract_text() for page in reader.pages]

    return "\n\n".join(page for page in pages if page and page.strip())


def _read_docx(path: Path) -> str:
    """Собирает текст из параграфов и таблиц DOCX."""
    doc = DocxDocument(path)

    # para.text и cell.text каждый раз заново обходят XML, поэтому их вычисление по одному разу
    parts = [text for para in doc.paragraphs if (text := para.text.strip())]

    # Текст в таблицах в doc.paragraphs не попадает - добор отдельно
    for table in doc.tables:
        for row in table.rows:
            cells = [text for cell in row.cells if (text := cell.text.strip())]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _read_plain_text(path: Path) -> str:
    """Читает текстовый файл, перебирая кодировки из TEXT_ENCODINGS."""
    data = path.read_bytes()

    for encoding in TEXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue

    # Последняя попытка: декодинг с заменой битых байтов, чтобы не потерять файл полностью
    print(f"Предупреждение: не удалось определить кодировку {path.name}, часть символов потеряна.")

    return data.decode("utf-8", errors="replace")


def load_text_from_file(file_path: Union[str, Path]) -> str:
    """
    Извлекает текст из файла формата PDF/DOCX/TXT/MD.
    Возвращает пустую строку, если формат не поддерживается или файл не читается.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"Файл не найден: {file_path}")

    extension = path.suffix.lower()

    try:
        match extension:
            case ".pdf":
                text = _read_pdf(path)
            case ".docx":
                text = _read_docx(path)
            case ".txt" | ".md":
                text = _read_plain_text(path)
            case _:
                print(f"Файл формата {extension} не поддерживается - файл пропущен.")
                return ""

    except Exception as e:
        print(f"Ошибка при чтении {path.name}: {e}")
        return ""

    return text.strip()


def load_documents_from_dir(directory: Union[str, Path]) -> List[LoadedDocument]:
    """
    Загружает все поддерживаемые документы из каталога и его подкаталогов.
    Файлы обрабатываются в отсортированном порядке, чтобы индекс был воспроизводимым.
    """
    root = Path(directory)

    if not root.is_dir():
        raise NotADirectoryError(f"Каталог не найден: {directory}")

    paths = sorted(
        path for path in root.glob("**/*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )

    if not paths:
        print(f"В каталоге {root} нет файлов поддерживаемых форматов ({', '.join(sorted(SUPPORTED_EXTENSIONS))}).")
        return []

    documents = []
    for path in paths:
        text = load_text_from_file(path)
        if not text:
            continue

        # Имя относительно каталога данных: различает файлы из вложенных папок
        source = path.relative_to(root).as_posix()
        documents.append(LoadedDocument(source=source, text=text))
        print(f"Загружен {source} ({len(text)} символов).")

    return documents
